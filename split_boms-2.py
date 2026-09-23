"""
split_boms.py — Separa un reporte consolidado de BOMs (Oracle EBS "Bill of
Material Structure Report") en un PDF por cada BOM.

Cada PDF de salida lleva:
  - La PORTADA del reporte (con el bloque de parámetros completo, que el
    encabezado de cada BOM no trae).
  - El contenido del BOM, desde su línea 'Organization:' hasta justo antes del
    siguiente BOM (el último, hasta el 'End of Report').
  - Un 'End of Report' al final.

Uso:
    1. Pon los reportes consolidados en la carpeta de entrada.
    2. Ajusta las rutas abajo si hace falta.
    3. python split_boms.py

Requiere:  pip install pymupdf openpyxl
"""

import re
import sys
from pathlib import Path

try:
    import fitz
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError as exc:
    print(f"Falta una dependencia: {exc}")
    print("Ejecuta: pip install pymupdf openpyxl")
    raise SystemExit(1)

__version__ = "1.2.0"

# ── CONFIGURACIÓN ───────────────────────────────────────────────────────────
CARPETA_ENTRADA = "Input_PDFs"       # reportes consolidados
CARPETA_SALIDA  = "Split_Output"     # un PDF por BOM
RESUMEN_XLSX    = "Split_Resumen.xlsx"

MARGEN_INFERIOR = 26.0    # margen útil al pie de página
SEPARACION      = 12.0    # espacio entre la portada y el contenido del BOM
# ────────────────────────────────────────────────────────────────────────────

# Códigos Oracle: 2-4 caracteres, guion, y el resto (CE-82-061Z-A, 4D-81-584Z-G,
# DWG-SA-004194-D, BOM-011538...)
ITEM_RE    = re.compile(r"\bItem:\s*([A-Z0-9]{2,4}-[A-Z0-9-]{4,})", re.I)
ORG_RE     = re.compile(r"\bOrganization:", re.I)
EOR_RE     = re.compile(r"End of Report", re.I)
FROM_TO_RE = re.compile(r"^(Item From:|To:)", re.I)
ITEM_LABEL_RE = re.compile(r"^Item:\s*$", re.I)
NOMBRE_MAL = re.compile(r'[<>:"/\\|?*]')


def nombre_seguro(valor):
    return NOMBRE_MAL.sub("_", str(valor).strip()).rstrip(". ") or "SIN_ITEM"


def _spans(page):
    for b in page.get_text("dict")["blocks"]:
        if b["type"] != 0:
            continue
        for line in b["lines"]:
            for s in line["spans"]:
                yield s


def lineas(page):
    """Devuelve [(y, texto, x0)] de la página, ordenado por y."""
    por_y = {}
    for s in _spans(page):
        por_y.setdefault(round(s["bbox"][1], 1), []).append(s)
    out = []
    for y in sorted(por_y):
        ss = sorted(por_y[y], key=lambda s: s["bbox"][0])
        out.append((y, "".join(s["text"] for s in ss),
                    min(s["bbox"][0] for s in ss)))
    return out


# ════════════════════════════════════════════════════════════════════════════
# LOCALIZAR LOS BOMs
# ════════════════════════════════════════════════════════════════════════════

def localizar_boms(doc):
    """Devuelve [{item, pagina, y_inicio}] — uno por BOM.

    y_inicio es la línea 'Organization:' que precede al 'Item:'; si no la hay,
    se usa la propia línea del Item.
    """
    encontrados = []
    for pn in range(len(doc)):
        ls = lineas(doc[pn])
        for i, (y, texto, _) in enumerate(ls):
            m = ITEM_RE.search(texto)
            if not m:
                continue
            # Buscar hacia atrás la línea 'Organization:' más cercana
            y_ini = y
            for j in range(i - 1, max(-1, i - 4), -1):
                if ORG_RE.search(ls[j][1]):
                    y_ini = ls[j][0]
                    break
            encontrados.append({
                "item": m.group(1).upper().rstrip(".,;:"),
                "pagina": pn,
                "y_inicio": y_ini,
            })

    # ── Quitar repeticiones del MISMO BOM ───────────────────────────────────
    # Cuando un BOM ocupa varias páginas, Oracle repite su encabezado (con el
    # mismo 'Item:') en cada una. Esas repeticiones NO son BOMs nuevos: se
    # descarta toda marca cuyo item sea igual al del BOM en curso.
    limpio = []
    for e in encontrados:
        if limpio and e["item"] == limpio[-1]["item"]:
            continue
        limpio.append(e)

    # Un mismo item podría reaparecer más adelante separado por otros BOMs
    # (poco habitual, pero entonces son entradas distintas del reporte). Se
    # avisa para que se pueda revisar.
    from collections import Counter
    repetidos = [it for it, n in Counter(e["item"] for e in limpio).items() if n > 1]
    if repetidos:
        print(f"   Aviso: {len(repetidos)} item(s) aparecen en bloques separados "
              f"del reporte: {', '.join(repetidos[:5])}"
              + (" ..." if len(repetidos) > 5 else ""))
    return limpio


def fin_del_reporte(doc):
    """(pagina, y) del 'End of Report' final."""
    for pn in range(len(doc) - 1, -1, -1):
        for y, texto, _ in lineas(doc[pn]):
            if EOR_RE.search(texto):
                return pn, y
    return len(doc) - 1, doc[len(doc) - 1].rect.height - MARGEN_INFERIOR


def limites(boms, doc):
    """Añade a cada BOM dónde termina: (pagina_fin, y_fin) = inicio del siguiente."""
    pn_eor, y_eor = fin_del_reporte(doc)
    for i, b in enumerate(boms):
        if i + 1 < len(boms):
            sig = boms[i + 1]
            b["pagina_fin"] = sig["pagina"]
            b["y_fin"] = sig["y_inicio"]
        else:
            b["pagina_fin"] = pn_eor
            b["y_fin"] = y_eor + 14      # incluir la línea del End of Report
    return boms


# ════════════════════════════════════════════════════════════════════════════
# CONSTRUIR EL PDF DE CADA BOM
# ════════════════════════════════════════════════════════════════════════════

def alto_portada(doc):
    """Hasta dónde llega el bloque de parámetros de la portada (pág. 1)."""
    ls = lineas(doc[0])
    ultimo = 0
    for y, texto, _ in ls:
        if ITEM_RE.search(texto) or ORG_RE.search(texto):
            break
        if texto.strip():
            ultimo = y
    return ultimo + 16


def ajustar_portada(pagina, item):
    """Personaliza la portada copiada para este BOM:
      - Borra el VALOR de 'Item From:' y 'To:' (son los extremos del rango con
        el que se descargó el reporte; no aplican a un BOM concreto).
      - Escribe el part number de ESTE BOM junto a la etiqueta 'Item:'.

    Se trabaja con las palabras reales de la página (no con posiciones
    estimadas) para no recortar las etiquetas ni las líneas vecinas.
    """
    tam = 9.75
    palabras = pagina.get_text("words")       # (x0, y0, x1, y1, texto, ...)

    # Agrupar por renglón
    por_y = {}
    for w in palabras:
        por_y.setdefault(round(w[1], 1), []).append(w)

    borrar = []       # rectángulos a limpiar
    escribir = []     # (x, y, texto) a insertar después

    for y in sorted(por_y):
        fila = sorted(por_y[y], key=lambda w: w[0])
        textos = [w[4] for w in fila]
        if not textos:
            continue
        primero = textos[0].strip().lower()

        # ── 'Item From:' / 'To:'  ->  borrar solo el valor ──────────────────
        es_item_from = (primero == "item" and len(textos) > 1
                        and textos[1].strip().lower().startswith("from"))
        es_to = primero == "to:"
        if es_item_from or es_to:
            # El valor son las palabras que siguen a la etiqueta (la que
            # termina en ':')
            idx_label = next((i for i, t in enumerate(textos) if t.endswith(":")), 0)
            valores = fila[idx_label + 1:]
            if valores:
                x_ini = min(w[0] for w in valores) - 1
                y0 = min(w[1] for w in valores)
                y1 = max(w[3] for w in valores)
                borrar.append(fitz.Rect(x_ini, y0, pagina.rect.width, y1))
            continue

        # ── 'Item:' sin valor  ->  poner el part number de este BOM ─────────
        if primero == "item:" and len(textos) == 1:
            w0 = fila[0]
            escribir.append((w0[2] + tam * 0.6, w0[1] + tam * 0.85, item))

    for r in borrar:
        pagina.add_redact_annot(r, fill=(1, 1, 1))
    if borrar:
        pagina.apply_redactions()
    for x, y, txt in escribir:
        pagina.insert_text(fitz.Point(x, y), txt, fontname="cour", fontsize=tam)


def construir_bom(doc, bom, portada_h):
    """Arma el PDF de un BOM: portada + su contenido."""
    w = doc[0].rect.width
    h = doc[0].rect.height
    salida = fitz.open()

    # ── Página 1: portada + inicio del contenido ────────────────────────────
    pagina = salida.new_page(width=w, height=h)
    pagina.show_pdf_page(fitz.Rect(0, 0, w, portada_h), doc, 0,
                         clip=fitz.Rect(0, 0, w, portada_h))
    ajustar_portada(pagina, bom["item"])

    y_cursor = portada_h + SEPARACION
    pn, y = bom["pagina"], bom["y_inicio"]
    pn_fin, y_fin = bom["pagina_fin"], bom["y_fin"]

    while pn <= pn_fin:
        desde = y if pn == bom["pagina"] else 0
        hasta = y_fin if pn == pn_fin else h - MARGEN_INFERIOR
        alto = hasta - desde
        if alto <= 2:
            pn += 1
            continue

        disponible = (h - MARGEN_INFERIOR) - y_cursor
        if alto > disponible:
            # Lo que cabe en esta página; el resto va en una nueva
            if disponible > 24:
                pagina.show_pdf_page(
                    fitz.Rect(0, y_cursor, w, y_cursor + disponible), doc, pn,
                    clip=fitz.Rect(0, desde, w, desde + disponible))
                desde += disponible
                alto -= disponible
            pagina = salida.new_page(width=w, height=h)
            y_cursor = MARGEN_INFERIOR
            # Puede necesitar varias páginas más
            while alto > (h - MARGEN_INFERIOR * 2):
                trozo = h - MARGEN_INFERIOR * 2
                pagina.show_pdf_page(
                    fitz.Rect(0, y_cursor, w, y_cursor + trozo), doc, pn,
                    clip=fitz.Rect(0, desde, w, desde + trozo))
                desde += trozo
                alto -= trozo
                pagina = salida.new_page(width=w, height=h)
                y_cursor = MARGEN_INFERIOR

        if alto > 2:
            pagina.show_pdf_page(
                fitz.Rect(0, y_cursor, w, y_cursor + alto), doc, pn,
                clip=fitz.Rect(0, desde, w, desde + alto))
            y_cursor += alto

        pn += 1
        y = 0

    # ── 'End of Report' propio ──────────────────────────────────────────────
    # Cada BOM separado debe cerrar con su propio marcador: es lo que delimita
    # la sección del BOM para las herramientas que lo procesan después.
    texto_eor = "***** End of Report *****"
    ya_tiene = any(EOR_RE.search(t) for _, t, _ in lineas(pagina))
    if not ya_tiene:
        if y_cursor > h - MARGEN_INFERIOR - 30:
            pagina = salida.new_page(width=w, height=h)
            y_cursor = MARGEN_INFERIOR
        tam = 9.75
        ancho = fitz.get_text_length(texto_eor, fontname="cour", fontsize=tam)
        pagina.insert_text(fitz.Point((w - ancho) / 2, y_cursor + 24 + tam * 0.85),
                           texto_eor, fontname="cour", fontsize=tam)

    return salida


# ════════════════════════════════════════════════════════════════════════════
# PROCESO
# ════════════════════════════════════════════════════════════════════════════

def procesar(pdf_path, carpeta_salida):
    doc = fitz.open(pdf_path)
    stem = Path(pdf_path).stem
    print(f"\n[{stem}]  {len(doc)} páginas")

    boms = limites(localizar_boms(doc), doc)
    print(f"   BOMs detectados: {len(boms)}")
    if not boms:
        print("   *** No se encontró ningún 'Item:' — revisa el formato ***")
        return []

    portada_h = alto_portada(doc)
    carpeta_salida.mkdir(parents=True, exist_ok=True)

    filas = []
    for i, b in enumerate(boms, 1):
        nombre = nombre_seguro(b["item"]) + ".pdf"
        destino = carpeta_salida / nombre
        n = 2
        while destino.exists():
            destino = carpeta_salida / f"{nombre_seguro(b['item'])}({n}).pdf"
            n += 1
        try:
            out = construir_bom(doc, b, portada_h)
            out.save(str(destino))
            out.close()
            filas.append({"item": b["item"], "archivo": destino.name,
                          "origen": stem, "pagina": b["pagina"] + 1,
                          "paginas_salida": "", "estado": "OK"})
            if i <= 5 or i % 25 == 0:
                print(f"   [{i:3d}/{len(boms)}] {b['item']} -> {destino.name}")
        except Exception as e:
            filas.append({"item": b["item"], "archivo": "", "origen": stem,
                          "pagina": b["pagina"] + 1, "paginas_salida": "",
                          "estado": f"ERROR: {e!r}"})
            print(f"   [{i:3d}] ERROR en {b['item']}: {e!r}")

    doc.close()
    return filas


def crear_resumen(filas, ruta):
    wb = Workbook(); ws = wb.active; ws.title = "Resumen"
    hdr = ["item", "archivo", "origen", "pagina_en_reporte", "estado"]
    fill = PatternFill("solid", fgColor="1F4E78")
    fnt = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    thin = Side(style="thin", color="D0D0D0")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    for j, h in enumerate(hdr, 1):
        c = ws.cell(1, j, h); c.fill = fill; c.font = fnt; c.border = bd
        c.alignment = Alignment(horizontal="center")
    for i, f in enumerate(filas, 2):
        for j, k in enumerate(["item", "archivo", "origen", "pagina", "estado"], 1):
            ws.cell(i, j, f.get(k, "")).border = bd
    for col, w in zip("ABCDE", [22, 34, 30, 18, 40]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(str(ruta))


def main():
    base = Path(__file__).parent
    entrada = base / CARPETA_ENTRADA
    salida = base / CARPETA_SALIDA

    print("=" * 62)
    print(f" Splitter de BOMs v{__version__}")
    print("=" * 62)

    if not entrada.exists():
        entrada.mkdir(parents=True, exist_ok=True)
        print(f"Se creó la carpeta '{CARPETA_ENTRADA}'. Pon ahí los reportes.")
        return

    # En Windows el sistema de archivos no distingue mayúsculas, así que
    # buscar "*.pdf" y "*.PDF" por separado devolvería CADA archivo DOS VECES.
    # Se deduplica por ruta absoluta en minúsculas.
    vistos, pdfs = set(), []
    for f in sorted(list(entrada.glob("*.pdf")) + list(entrada.glob("*.PDF"))):
        clave = str(f.resolve()).lower()
        if clave not in vistos:
            vistos.add(clave)
            pdfs.append(f)
    if not pdfs:
        print(f"No hay PDFs en '{CARPETA_ENTRADA}'.")
        return

    todas = []
    for p in pdfs:
        todas.extend(procesar(p, salida))

    crear_resumen(todas, base / RESUMEN_XLSX)
    ok = sum(1 for f in todas if f["estado"] == "OK")
    print("\n" + "=" * 62)
    print(f" BOMs separados OK : {ok}")
    print(f" Con error         : {len(todas) - ok}")
    print(f" Carpeta de salida : {salida}")
    print(f" Resumen           : {RESUMEN_XLSX}")
    print("=" * 62)


if __name__ == "__main__":
    main()
