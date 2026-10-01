"""
renombrar_boms.py — Agrega el código del BOM y la revisión al nombre de los PDFs.

    4D-81-584Z-G.pdf   ->   4D-81-584Z-G(BOM-005379)_REV001.pdf

El código del BOM se toma de un Excel que relaciona cada archivo con su código.
La revisión se toma del Excel si se indica; si no, se lee de la fila del BOM
dentro del PDF; si tampoco está, se usa 001.

Uso:
    1. Ajusta las rutas en CONFIGURACIÓN.
    2. Deja SIMULAR = True para ver qué haría sin tocar nada.
    3. Cuando el resultado te convenza, pon SIMULAR = False y vuelve a correrlo.

Requiere:  pip install pymupdf openpyxl
"""

import os
import re
import shutil
from pathlib import Path

try:
    import fitz
    import openpyxl
except ImportError as exc:
    print(f"Falta una dependencia: {exc}")
    print("Ejecuta: pip install pymupdf openpyxl")
    raise SystemExit(1)

__version__ = "1.0.0"

# ── CONFIGURACIÓN ───────────────────────────────────────────────────────────
CARPETA   = "PDFs"                  # carpeta con los PDFs a renombrar
EXCEL     = "codigos_bom.xlsx"      # Excel: nombre_archivo | codigo_bom | rev
SIMULAR   = True                    # True = solo muestra; False = renombra
COPIAR_A  = ""                      # vacío = renombrar en sitio.
                                    # Con una ruta, copia los renombrados ahí
                                    # y deja los originales intactos.
# ────────────────────────────────────────────────────────────────────────────

NOMBRE_MAL = re.compile(r'[<>:"/\\|?*]')
BOM_RE     = re.compile(r"\bBOM-\d{3,}\b", re.I)


def nombre_seguro(valor):
    return NOMBRE_MAL.sub("_", str(valor).strip()).rstrip(". ")


def fmt_rev(valor):
    """Formatea la revisión a 3 dígitos: 9 -> '009', '10' -> '010'."""
    try:
        return f"{int(str(valor).strip()):03d}"
    except (TypeError, ValueError):
        return str(valor).strip()


# ════════════════════════════════════════════════════════════════════════════
# LEER EL EXCEL
# ════════════════════════════════════════════════════════════════════════════

# Nombres aceptados para cada columna (por si el encabezado varía)
ALIAS = {
    "archivo":  {"nombre_archivo", "archivo", "item", "nombre", "pdf",
                 "nombre oracle", "nombre_oracle"},
    "bom":      {"codigo_bom", "bom", "bom_code", "new_bom", "codigo"},
    "rev":      {"rev", "revision", "rev_bom", "revision_bom"},
}


def leer_excel(ruta):
    """Devuelve {nombre_archivo_en_minusculas: (codigo_bom, rev_o_None)}."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    mapa = {}

    for ws in wb.worksheets:
        # Buscar la fila de encabezados (puede haber una nota arriba)
        hdr_i, col = None, {}
        for i, fila in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True)):
            celdas = [str(c).strip().lower() if c is not None else "" for c in fila]
            for j, h in enumerate(celdas):
                for destino, nombres in ALIAS.items():
                    if h in nombres and destino not in col:
                        col[destino] = j
            if "archivo" in col and "bom" in col:
                hdr_i = i
                break
            col = {}
        if hdr_i is None:
            continue

        for fila in ws.iter_rows(min_row=hdr_i + 2, values_only=True):
            if not any(c is not None and str(c).strip() for c in fila):
                continue
            archivo = fila[col["archivo"]] if col["archivo"] < len(fila) else None
            bom     = fila[col["bom"]]     if col["bom"]     < len(fila) else None
            rev     = (fila[col["rev"]] if "rev" in col and col["rev"] < len(fila)
                       else None)
            if not archivo or not bom:
                continue
            clave = os.path.splitext(str(archivo).strip())[0].lower()
            mapa[clave] = (str(bom).strip(), str(rev).strip() if rev else None)
        if mapa:
            break      # la primera hoja útil basta

    return mapa


# ════════════════════════════════════════════════════════════════════════════
# LEER LA REVISIÓN DEL PDF
# ════════════════════════════════════════════════════════════════════════════

# Unidades de medida habituales. La columna Rev va justo ANTES de la UOM, y
# esa es la forma fiable de localizarla: buscar "el primer valor corto" falla
# porque la descripción (p. ej. "BOM NUMBER FOR") está antes y también es corta.
UOM = {"EA", "ML", "GR", "BX", "KG", "LB", "FT", "IN", "CM", "MM", "L", "M",
       "OZ", "PK", "RL", "SH", "YD", "CS", "DZ", "GAL", "BG"}


def rev_desde_pdf(pdf_path, codigo_bom):
    """Busca la fila del BOM dentro del PDF y devuelve el valor de su columna
    Rev: la palabra inmediatamente anterior a la UOM de esa fila.
    Devuelve None si el BOM no aparece como componente."""
    try:
        doc = fitz.open(pdf_path)
    except Exception:
        return None

    base = codigo_bom.upper().split("_R")[0]
    try:
        for pn in range(len(doc)):
            palabras = doc[pn].get_text("words")      # (x0, y0, x1, y1, texto…)
            for x0, y0, x1, y1, t, *_ in palabras:
                if not t.upper().startswith(base):
                    continue
                fila = sorted(
                    (w for w in palabras if abs(w[1] - y0) < 4 and w[0] > x1),
                    key=lambda w: w[0])
                textos = [w[4].strip() for w in fila]
                for i, v in enumerate(textos):
                    if v.upper() in UOM and i > 0:
                        anterior = textos[i - 1]
                        if 1 <= len(anterior) <= 4 and (anterior.isdigit()
                                                        or anterior.isalpha()):
                            return anterior
                        return None
                return None
    finally:
        doc.close()
    return None


# ════════════════════════════════════════════════════════════════════════════
# PROCESO
# ════════════════════════════════════════════════════════════════════════════

REV_SUFIJO_RE = re.compile(r"[_\-\s]*REV[\s_\-]*[A-Z0-9]{1,4}\s*$", re.I)


def nuevo_nombre(stem, codigo_bom, rev):
    """item(BOM-XXXX)_REVxxx

    Es idempotente: si el nombre ya trae el código del BOM no lo repite, y si
    ya termina en _REVxxx lo sustituye en vez de encadenar otro sufijo.
    """
    stem = REV_SUFIJO_RE.sub("", stem).rstrip(" _-")
    if codigo_bom.lower() in stem.lower():
        return f"{nombre_seguro(stem)}_REV{rev}.pdf"
    return f"{nombre_seguro(stem)}({nombre_seguro(codigo_bom)})_REV{rev}.pdf"


def main():
    base = Path(__file__).parent
    carpeta = base / CARPETA
    excel = base / EXCEL

    print("=" * 64)
    print(f" Renombrar BOMs v{__version__}" + ("   [SIMULACIÓN]" if SIMULAR else ""))
    print("=" * 64)

    if not carpeta.is_dir():
        print(f"No existe la carpeta '{CARPETA}'.")
        return
    if not excel.is_file():
        print(f"No existe el Excel '{EXCEL}'.")
        return

    mapa = leer_excel(excel)
    print(f"Excel: {len(mapa)} archivo(s) con código de BOM")
    if not mapa:
        print("  *** No se encontraron columnas 'nombre_archivo' y 'codigo_bom' ***")
        return

    # Deduplicar: en Windows *.pdf y *.PDF devuelven el mismo archivo
    vistos, pdfs = set(), []
    for f in sorted(list(carpeta.glob("*.pdf")) + list(carpeta.glob("*.PDF"))):
        clave = str(f.resolve()).lower()
        if clave not in vistos:
            vistos.add(clave)
            pdfs.append(f)
    print(f"PDFs en la carpeta: {len(pdfs)}\n")

    destino_dir = (base / COPIAR_A) if COPIAR_A else carpeta
    if COPIAR_A and not SIMULAR:
        destino_dir.mkdir(parents=True, exist_ok=True)

    hechos = sin_codigo = errores = 0

    for pdf in pdfs:
        stem = pdf.stem
        entrada = mapa.get(stem.lower())

        # Si no hay match exacto, probar que el nombre CONTENGA la clave
        if entrada is None:
            for clave, val in mapa.items():
                if clave in stem.lower():
                    entrada = val
                    break

        if entrada is None:
            print(f"  [—] {pdf.name}: sin código en el Excel (se deja igual)")
            sin_codigo += 1
            continue

        codigo_bom, rev_excel = entrada

        # Revisión: la del Excel; si no, la del PDF; si no, 001
        if rev_excel:
            rev, origen = fmt_rev(rev_excel), "Excel"
        else:
            leida = rev_desde_pdf(pdf, codigo_bom)
            if leida:
                rev, origen = fmt_rev(leida), "PDF"
            else:
                rev, origen = "001", "por defecto"

        nombre = nuevo_nombre(stem, codigo_bom, rev)
        destino = destino_dir / nombre

        if destino.resolve() == pdf.resolve():
            print(f"  [=] {pdf.name}: ya tiene ese nombre")
            continue

        n = 2
        while destino.exists():
            destino = destino_dir / nombre.replace(".pdf", f"({n}).pdf")
            n += 1

        print(f"  {pdf.name}")
        print(f"      -> {destino.name}   (rev desde {origen})")

        if not SIMULAR:
            try:
                if COPIAR_A:
                    shutil.copy2(pdf, destino)
                else:
                    pdf.rename(destino)
                hechos += 1
            except Exception as e:
                print(f"      *** ERROR: {e!r}")
                errores += 1
        else:
            hechos += 1

    print("\n" + "=" * 64)
    verbo = "Se renombrarían" if SIMULAR else ("Copiados" if COPIAR_A else "Renombrados")
    print(f" {verbo}        : {hechos}")
    print(f" Sin código en Excel: {sin_codigo}")
    if errores:
        print(f" Con error          : {errores}")
    if SIMULAR:
        print("\n Esto fue una SIMULACIÓN: no se tocó ningún archivo.")
        print(" Pon SIMULAR = False arriba para aplicarlo de verdad.")
    print("=" * 64)


if __name__ == "__main__":
    main()
